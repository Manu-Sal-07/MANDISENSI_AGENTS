"""
Deep-dive study guide for the Review-2 deck: for every slide, the idea in plain words, the technical content explained, the numbers and what they mean,
why it was done that way, the honest limits and likely questions.

    python report/build_deep_dive.py   ->  report/MandiSense_Review2_Deep_Dive_Guide.docx

Content lives in dd_part1.py, dd_part2.py, dd_part3.py. Styling helpers are shared with build_deck_script.py.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import dd_part1 as P1  # noqa: E402
import dd_part2 as P2  # noqa: E402
import dd_part3 as P3  # noqa: E402
from build_deck_script import BLUE, NAVY, para, rich, shade, table as _table  # noqa: E402

OUT = Path(__file__).resolve().parent / "MandiSense_Review2_Deep_Dive_Guide.docx"
import re

_EM = re.compile(r"(?<!\*)\*(?!\*)([^*\n]+?)(?<!\*)\*(?!\*)")


def clean(x):
    """Single-asterisk emphasis is not rendered by rich(); drop the markers."""
    if isinstance(x, str):
        return _EM.sub(r"\1", x)
    if isinstance(x, dict):
        return {k: clean(v) for k, v in x.items()}
    if isinstance(x, list):
        return [clean(v) for v in x]
    if isinstance(x, tuple):
        return tuple(clean(v) for v in x)
    return x


SLIDES = clean(sorted(P1.SLIDES + P2.SLIDES + P3.SLIDES, key=lambda s: s["n"]))
P1.FRONT = clean(P1.FRONT)
P3.APPX = clean(P3.APPX)


def table(d, rows, widths, header=True, size=9.5):
    t = _table(d, rows, widths, header, size)
    t.autofit = False
    for j, w in enumerate(widths):
        t.columns[j].width = Cm(w)
        for c in t.columns[j].cells:
            c.width = Cm(w)
    return t


def label(d, text, color=NAVY):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = color
    return p


def callout(d, text, fill="EAF2F8"):
    t = d.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, fill)
    c.text = ""
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    rich(p, text, 10.5)
    d.add_paragraph().paragraph_format.space_after = Pt(0)


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MandiSense AI · Review 2 deep-dive guide · page ")
    r.font.size = Pt(8.5)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        run = p.add_run()
        run.font.size = Pt(8.5)
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)


def bullets(d, items, space=3):
    for it in items:
        para(d, it, style="List Bullet", space=space)


def slide_block(d, s):
    h = d.add_heading(f"Slide {s['n']} · {s['title']}", level=1)
    h.paragraph_format.space_before = Pt(22)
    callout(d, f"**In one line.** {s['one']}")
    if s.get("plain"):
        label(d, "In plain words")
        for t in s["plain"]:
            para(d, t, space=4)
    if s.get("deeper"):
        label(d, "Go deeper: the technical content, explained")
        bullets(d, s["deeper"])
    if s.get("numbers"):
        label(d, "The numbers on this slide and what each one means")
        table(d, [["Number", "What it means"]] + [[a, b] for a, b in s["numbers"]], [5.2, 11.8], size=9.5)
    if s.get("why"):
        label(d, "Why we did it this way")
        bullets(d, s["why"])
    if s.get("limits"):
        label(d, "Limits and things to fix (say these yourself)", RGBColor(0x9C, 0x3B, 0x00))
        bullets(d, s["limits"])
    if s.get("ask"):
        label(d, "If asked")
        for q, a in s["ask"]:
            para(d, f"**Q. {q}**", space=1)
            para(d, a, space=5)


def main() -> None:
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(1.9)
    sec.top_margin = sec.bottom_margin = Cm(1.7)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        h = d.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = NAVY
    footer(sec)

    t = d.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("MandiSense AI: Review 2\nDeep-Dive Study Guide, slide by slide")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    para(d, "Every idea, formula and number on every slide, explained in plain words with the technical content kept. "
            "Companion to the speaker notes (which stay as they are).", 11, italic=True, space=4).alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(d, "B.M.S. College of Engineering · Project Work Phase 2 (23DS7PWPP2) · Srijan C Vachadmath, Manohara Salmani, Shreejith Shah · Guide: Dr. Kalyan N", 9.5, italic=True, space=10).alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading("How to use this guide", level=1)
    bullets(d, P1.FRONT["how"])

    d.add_heading("Slide index", level=2)
    rows = [["Slide", "Title", "One line"]] + [[str(s["n"]), s["title"], s["one"]] for s in SLIDES]
    table(d, rows, [1.4, 5.6, 10.0], size=8.5)

    d.add_heading("The project in five paragraphs", level=1)
    for line in P1.FRONT["storyline"]:
        para(d, line, style="List Number", space=3)

    d.add_heading("Numbers worth memorising", level=1)
    table(d, [["Number", "What it is"]] + [[a, b] for a, b in P1.FRONT["numbers"]], [5.6, 11.4], size=9.5)

    d.add_heading("Part A. Glossary: every technical word, once, in plain language", level=1)
    para(d, "Learn these once and the slides become readable. When speaking, use the plain explanation first and the technical word second.", italic=True, space=3)
    table(d, [["Term", "Plain meaning"]] + [[a, b] for a, b in P1.FRONT["glossary"]], [5.0, 12.0], size=9.5)
    d.add_page_break()

    d.add_heading("Part B. Slide by slide", level=1)
    para(d, "Slides 14 to 27 (setup, results, validation, benchmarking, ablation) are the hardest; read those entries twice.", italic=True, space=4)
    d.add_page_break()
    for s in SLIDES:
        slide_block(d, s)

    d.add_heading("Part C. Formula cheat sheet", level=1)
    para(d, "Every formula in the deck, in one line each, with what it means.", italic=True, space=3)
    table(d, [["Name", "Formula", "In plain words"]] + [list(r) for r in P3.APPX["formulas"]], [3.0, 5.6, 8.4], size=9)

    d.add_heading("Part D. Numbers that look contradictory but are not", level=1)
    table(d, [["Looks odd", "Explanation"]] + [list(r) for r in P3.APPX["traps"]], [5.6, 11.4], size=9.5)

    d.add_heading("Part E. Fix these on the slides before presenting", level=1)
    bullets(d, P3.APPX["fixes"])

    d.add_heading("Part F. Question bank by theme", level=1)
    para(d, "Short, honest answers. Say the first sentence and stop.", italic=True, space=3)
    theme = ""
    for th, q, a in P3.APPX["qa"]:
        if th != theme:
            label(d, th)
            theme = th
        para(d, f"**Q. {q}**", space=1)
        para(d, a, space=5)

    d.save(OUT)
    print("saved", OUT, "slides:", len(SLIDES))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
